#!/usr/bin/env python3
"""Build ТЕСТ_ПТС_вводное.docx: 6 questions + answers on page 2."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

OUT = Path(__file__).resolve().parents[1] / "ТЕСТ_ПТС_вводное.docx"


def set_run_font(run, size=14, bold=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_q(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(8)
    set_run_font(p.add_run(f"{num}. {text}"), 14)


def add_ans(doc, title, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(title), 14, bold=True)
    for line in lines:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(1)
        p2.paragraph_format.left_indent = Pt(12)
        set_run_font(p2.add_run(line), 12)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Emu(7560310)
    sec.page_height = Emu(10692130)
    sec.left_margin = Emu(1080135)
    sec.right_margin = Emu(539750)
    sec.top_margin = Emu(720090)
    sec.bottom_margin = Emu(720090)

    h = doc.add_paragraph()
    r = h.add_run("ТЕСТ_ТЕХ · ПТС · вводное занятие")
    set_run_font(r, 14, bold=True)

    add_q(doc, 1, "Соедините термин и его задачи:")
    doc.add_paragraph()

    rows = [
        ("Выпускающий режиссёр", "А) строит кадр, держит фокус, слушает команды"),
        ("Режиссёр трансляции", "Б) связь с организаторами по сценарию"),
        ("Инженер ПТС", "В) отвечает за эфир целиком"),
        ("Оператор", "Г) графика в эфире"),
        ("Виджей трансляции", "Д) решает, что в программе: переключения, темп, крупности"),
        ("Продюсер", "Е) коммутация, сигналы, железо"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "775")
    tbl_ind.set(qn("w:type"), "dxa")
    table._tbl.tblPr.append(tbl_ind)
    for i, (left, right) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text = ""
        c1.text = ""
        set_run_font(c0.paragraphs[0].add_run(left), 12)
        set_run_font(c1.paragraphs[0].add_run(right), 12)

    add_q(doc, 2, "Зачем нужны трансляции? Дайте определение «ПТС».")
    add_q(doc, 3, "Какими бывают трансляции? Назовите форматы")
    add_q(doc, 4, "Опишите пайплайн трансляции от источника до зрителя.")
    add_q(doc, 5, "Что относится к железу, а что к софту на трансляции?")
    add_q(doc, 6, "Куда может уходить картинка с мероприятия? Приведите примеры.")

    pb = doc.add_paragraph()
    pb.add_run().add_break(WD_BREAK.PAGE)

    ah = doc.add_paragraph()
    set_run_font(ah.add_run("Ответы (для быстрой проверки)"), 14, bold=True)

    add_ans(
        doc,
        "1.",
        [
            "Выпускающий режиссёр — В",
            "Режиссёр трансляции — Д",
            "Инженер ПТС — Е",
            "Оператор — А",
            "Виджей трансляции — Г",
            "Продюсер — Б",
        ],
    )
    add_ans(
        doc,
        "2.",
        [
            "Зачем: доступность; расширение аудитории; удалённые команды и конференции; качество события.",
            "ПТС — передвижная телевизионная станция; цепочка, которая увозит событие за пределы зала (направление ТЕХ ГУАП).",
        ],
    )
    add_ans(doc, "3.", ["Форматы: Live; гибрид; запись / IMAG; студийная."])
    add_ans(doc, "4.", ["Источник → сигнал → микшер → доставка → зритель."])
    add_ans(
        doc,
        "5.",
        [
            "Железо: камеры, коммутация, микшер.",
            "Софт: vMix, контент (титры, бамперы, стингеры), кодеки и доставка, сети и стандарты (NDI, Dante, RTMP/SRT).",
        ],
    )
    add_ans(
        doc,
        "6.",
        [
            "Прямой эфир (VK, YouTube и др.); онлайн-встреча (Zoom, TrueConf и др.); запись (архив, нарезки, отчётность); удалённый зал.",
        ],
    )

    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
